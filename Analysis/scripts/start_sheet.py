"""Add a plain 'Start here' sheet to a deliverable workbook, so a reader knows what to open first."""
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink

NAVY = "001F5F"
BLUE = "1C5CAB"


def add_start_sheet(path, title, what, how, sheets):
    """what: one sentence on the workbook; how: list of short reading steps; sheets: [(sheet name, what it shows)]."""
    wb = load_workbook(path, rich_text=True)  # keep bold labels inside cells
    if "Start here" in wb.sheetnames:
        del wb["Start here"]
    ws = wb.create_sheet("Start here", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A3"] = what
    ws["A3"].alignment = wrap
    ws.merge_cells("A3:B3")
    ws.row_dimensions[3].height = 34

    ws["A5"] = "How to read it"
    ws["A5"].font = Font(bold=True, size=12, color=NAVY)
    r = 6
    for i, step in enumerate(how, 1):
        ws.cell(r, 1, f"Step {i}").font = Font(bold=True)
        ws.cell(r, 1).alignment = wrap
        ws.cell(r, 2, step).alignment = wrap
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    ws.cell(r, 1, "Sheet (click to open)").font = Font(bold=True, color="FFFFFF")
    ws.cell(r, 2, "What it shows").font = Font(bold=True, color="FFFFFF")
    for c in (1, 2):
        ws.cell(r, c).fill = PatternFill("solid", fgColor=BLUE)
    edge = Side(style="thin", color="C9C9C9")
    for name, desc in sheets:
        r += 1
        a = ws.cell(r, 1, name)
        a.hyperlink = Hyperlink(ref=a.coordinate, location=f"'{name}'!A1", display=name)
        a.font = Font(color="0563C1", underline="single")
        b = ws.cell(r, 2, desc)
        b.alignment = wrap
        for c in (a, b):
            c.border = Border(edge, edge, edge, edge)
            c.alignment = wrap
        ws.row_dimensions[r].height = 30
    # tabs in reading order: Start here, then the listed sheets, then any others
    listed = [wb[n] for n, _ in sheets if n in wb.sheetnames]
    rest = [s for s in wb.worksheets if s not in listed and s.title != "Start here"]
    wb._sheets = [ws] + listed + rest
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0
    for other in wb.worksheets[1:]:
        other.sheet_view.tabSelected = False
    wb.save(path)
